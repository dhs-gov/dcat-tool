"""Regression coverage for importing workbooks with multiple inventory sheets."""

import sys
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
import pytest

sys.path.append(str(Path(__file__).resolve().parents[1]))

from template_reader import TemplateReader


@pytest.mark.parametrize("second_properties", [
    ["dct:identifier", "dct:title"],
    ["dct:title", "dct:identifier"],
    ["dct:identifier", "dct:title", "dct:description"],
    ["dct:identifier"],
])
@pytest.mark.parametrize("first_has_data", [True, False])
def test_inventory_records_use_each_worksheets_header(
        tmp_path, second_properties, first_has_data):
    workbook = openpyxl.Workbook()
    workbook.active.title = "Instructions"
    workbook.active.append(["Do not import this row"])
    records = [
        {"dct:identifier": "dataset-1", "dct:title": "First dataset"},
        {"dct:identifier": "dataset-2", "dct:title": "Second dataset",
         "dct:description": "Second description"},
    ]
    expected = []
    for index, properties in enumerate([
            ["dct:identifier", "dct:title"], second_properties]):
        sheet = workbook.create_sheet("Inventory {}".format(index + 1))
        for column, name in enumerate(properties, 1):
            cell = sheet.cell(row=1, column=column, value=name)
            cell.comment = Comment("<{}>".format(name), "")
        sheet.append([None] * len(properties))
        if index == 1 or first_has_data:
            sheet.append([records[index][name] for name in properties])
            expected.append({name: records[index][name] for name in properties})

    filename = tmp_path / "inventories.xlsx"
    workbook.save(filename)
    workbook.close()
    reader = TemplateReader(filename)
    try:
        assert reader.inventory_records() == expected
    finally:
        reader.wb.close()
